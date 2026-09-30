"""Build the living Word report for the Adaptive Chess Engine project."""

from __future__ import annotations

from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Adaptive_Chess_Engine_Project_Report.docx"
ASSET_DIR = ROOT / "docs" / "report_assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F3F7FB"
GRAY = "666666"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(value))
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[i])
            if row_index % 2:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for number, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.22)
        p.paragraph_format.first_line_indent = Inches(-.22)
        p.add_run(f"{number}. ").bold = True
        p.add_run(item)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(89, 89, 89)


def add_figure(doc, path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    add_caption(doc, caption)


def add_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p_pr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), LIGHT_GRAY)
    pbdr.append(bottom)
    p_pr.append(pbdr)


def add_status_line(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(label + " ")
    r.bold = True
    r.font.color.rgb = RGBColor(23, 54, 93)
    p.add_run(text)


def _fonts():
    regular = Path(r"C:\Windows\Fonts\arial.ttf")
    bold = Path(r"C:\Windows\Fonts\arialbd.ttf")
    return (
        ImageFont.truetype(str(regular), 28),
        ImageFont.truetype(str(bold), 29),
        ImageFont.truetype(str(regular), 23),
    )


def _centered(draw, xy, text, font, fill="#222222", anchor="mm"):
    draw.multiline_text(xy, text, font=font, fill=fill, anchor=anchor, align="center", spacing=5)


def _arrow(draw, start, end, fill="#555555", width=4):
    draw.line([start, end], fill=fill, width=width)
    x, y = end
    draw.polygon([(x, y), (x-14, y-8), (x-14, y+8)], fill=fill)


def make_charts():
    font, bold, small = _fonts()
    size = (1400, 650)

    elo = ASSET_DIR / "elo_progression.png"
    im = Image.new("RGB", size, "white"); d = ImageDraw.Draw(im)
    left, top, bottom, right = 180, 70, 540, 1320
    for val in range(1750, 2101, 50):
        y = bottom - (val-1750)/(2100-1750)*(bottom-top)
        d.line((left,y,right,y), fill="#E5E5E5", width=2)
        d.text((left-20,y), str(val), font=small, fill="#555", anchor="rm")
    points = [(470, bottom-(1871-1750)/350*(bottom-top)), (1040, bottom-(1996-1750)/350*(bottom-top))]
    d.line(points, fill="#2F75B5", width=8)
    for (x,y), val, low, high, label in zip(points,[1871,1996],[1792,1930],[1949,2063],["Earlier baseline\n78 games","Optimized baseline\n112 games"]):
        y1=bottom-(low-1750)/350*(bottom-top); y2=bottom-(high-1750)/350*(bottom-top)
        d.line((x,y1,x,y2), fill="#333", width=5); d.line((x-22,y1,x+22,y1), fill="#333", width=5); d.line((x-22,y2,x+22,y2), fill="#333", width=5)
        d.ellipse((x-13,y-13,x+13,y+13), fill="#2F75B5")
        d.text((x,y-35),str(val),font=bold,fill="#17365D",anchor="ms")
        _centered(d,(x,595),label,font)
    d.text((left,25),"Stockfish limited-strength scale",font=small,fill="#333",anchor="lm")
    im.save(elo)

    drift = ASSET_DIR / "rating_drift.png"
    im = Image.new("RGB", size, "white"); d = ImageDraw.Draw(im)
    left, top, bottom = 170, 55, 550
    for val in range(400,1401,200):
        y=bottom-(val-300)/1100*(bottom-top); d.line((left,y,1340,y),fill="#E5E5E5",width=2); d.text((left-20,y),str(val),font=small,fill="#555",anchor="rm")
    xs=[390,750,1110]; means=[772.5,985.1,1186.1]; lows=[370,897,1056]; highs=[950,1095,1296]; colors=["#9DC3E6","#5B9BD5","#17365D"]
    for x,mean,low,high,label,color in zip(xs,means,lows,highs,["Train","Validation","Test"],colors):
        y=bottom-(mean-300)/1100*(bottom-top); ylo=bottom-(low-300)/1100*(bottom-top); yhi=bottom-(high-300)/1100*(bottom-top)
        d.rectangle((x-95,y,x+95,bottom),fill=color); d.line((x,ylo,x,yhi),fill="#333",width=5); d.line((x-20,ylo,x+20,ylo),fill="#333",width=5); d.line((x-20,yhi,x+20,yhi),fill="#333",width=5)
        d.text((x,y-18),f"mean {mean:.1f}",font=bold,fill="#222",anchor="ms"); d.text((x,600),label,font=font,fill="#222",anchor="mm")
    d.text((left,25),"Chess.com rating",font=small,fill="#333",anchor="lm")
    im.save(drift)

    models = ASSET_DIR / "model_log_loss.png"
    im = Image.new("RGB", size, "white"); d = ImageDraw.Draw(im)
    names=["Constant","Logistic","Hierarchical","Gradient boosting"]; vals=[.421234,.404990,.404707,.393484]; colors=["#BFBFBF","#9DC3E6","#5B9BD5","#17365D"]
    x0,x1=370,1260; minv,maxv=.385,.425
    for i,(name,val,color) in enumerate(zip(names,vals,colors)):
        y=90+i*125; width=(val-minv)/(maxv-minv)*(x1-x0)
        d.text((x0-25,y+34),name,font=font,fill="#222",anchor="rm"); d.rectangle((x0,y,x0+width,y+68),fill=color); d.text((x0+width+15,y+34),f"{val:.5f}",font=small,fill="#222",anchor="lm")
    d.text((815,610),"Test log loss   lower is better",font=font,fill="#333",anchor="mm")
    im.save(models)

    architecture = ASSET_DIR / "architecture.png"
    im = Image.new("RGB", (1500,720), "white"); d=ImageDraw.Draw(im)
    boxes=[(40,245,275,440,"Chess GUI\nUCI client","#D9EAF7"),(375,120,790,610,"C++ production engine\n\nRules and position\nEvaluation and search\nMultiPV and adaptive selector","#B4C7E7"),(900,80,1460,305,"Python validation layer\nPerft  tests  oracle  benchmarks","#E2F0D9"),(900,400,1460,655,"Python research layer\nPGN extraction  labels  features\nmodel training  match analysis","#FFF2CC")]
    for x1,y1,x2,y2,textv,fill in boxes:
        d.rounded_rectangle((x1,y1,x2,y2),radius=22,fill=fill,outline="#777",width=3); _centered(d,((x1+x2)/2,(y1+y2)/2),textv,font)
    _arrow(d,(275,342),(375,342)); d.text((325,315),"UCI",font=small,fill="#555",anchor="mm")
    _arrow(d,(790,210),(900,192)); d.text((845,165),"cross-check",font=small,fill="#555",anchor="mm")
    _arrow(d,(900,520),(790,470)); d.text((845,545),"profiles and evidence",font=small,fill="#555",anchor="mm")
    im.save(architecture)

    chain = ASSET_DIR / "research_chain.png"
    im=Image.new("RGB",(1600,430),"white"); d=ImageDraw.Draw(im)
    labels=["Chess.com\nPGNs","Chronological\ndecisions","Stockfish\nlabels","Position\nfeatures","Three\nmodels","Held-out\nevidence","C++ selector\nand matches"]; fills=["#D9EAF7","#D9EAF7","#FFF2CC","#E2F0D9","#E2F0D9","#E2F0D9","#F4CCCC"]
    for i,(label,fill) in enumerate(zip(labels,fills)):
        x1=20+i*225; x2=x1+185; d.rounded_rectangle((x1,130,x2,300),radius=18,fill=fill,outline="#777",width=3); _centered(d,((x1+x2)/2,215),label,small)
        if i<6:_arrow(d,(x2,215),(x2+35,215),width=3)
    im.save(chain)
    return elo, drift, models, architecture, chain


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(.72)
    section.bottom_margin = Inches(.68)
    section.left_margin = Inches(.82)
    section.right_margin = Inches(.82)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(35, 35, 35)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, before, after in [
        ("Title", 27, 0, 12), ("Subtitle", 12, 0, 12),
        ("Heading 1", 20, 14, 7), ("Heading 2", 14, 10, 5), ("Heading 3", 11, 8, 4)
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_ppr = styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(10)
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "Adaptive Chess Engine  Living Project Report"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.runs[0].font.size = Pt(8)
    p.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(footer.add_paragraph())


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def build_report():
    elo, drift, models, architecture, chain = make_charts()
    doc = Document()
    configure_document(doc)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    r = p.add_run("ADAPTIVE CHESS ENGINE")
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(47,117,181)
    title = doc.add_paragraph(style="Title")
    title.add_run("Project Development and Research Report")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("A living technical record from engine foundations to personalized opponent modelling")
    add_status_line(doc, "Project owner", "Yash Dutt")
    add_status_line(doc, "Repository", "Adaptive Chess Engine")
    add_status_line(doc, "Report version", "1.0  30 September 2026")
    add_status_line(doc, "Repository history reviewed", "All project commits through f3cd286 on main")
    doc.add_paragraph()
    para(doc, "This report is the durable context record for the project. It explains what was built, why design choices were made, how claims were measured, what the results mean, and what remains unproven. It is intentionally more complete than a release note and more accessible than the source code.")
    para(doc, "The project has reached two distinct achievements. First, a correct Python engine was developed and transformed into a measured C++ production engine whose local short-time benchmark improved from an estimated 1871 to 1996 Elo on Stockfish 18's limited-strength scale. Second, a bounded adaptive layer and a reproducible opponent-modelling pipeline were built. The current evidence shows that a nonlinear personalized model predicts future YashDutt7 errors better than a same-class population model, but the project has not yet shown that this predictive advantage produces more wins at equal compute.")
    doc.add_page_break()

    doc.add_heading("How to use this report", level=1)
    para(doc, "Read Part One to understand the conventional engine and the performance foundation that makes adaptive experiments credible. Read Part Two to understand the adaptive selector, personal-game dataset, model comparison, limitations, and next experiment. The final pages provide a milestone ledger and a maintenance rule so future work is not separated from its evidence.")
    doc.add_heading("Contents", level=2)
    contents = [
        "Executive summary", "Part One  Main performance development", "1  Foundations and correctness", "2  Search and evaluation", "3  Performance engineering", "4  C++ production migration", "5  Measurement and current engine baseline",
        "Part Two  Adaptive layer and opponent modelling", "6  Research question and safety contract", "7  Hand-authored adaptive layer", "8  Personal-data research pipeline", "9  Three-model comparison", "10  Player growth and concept drift", "11  What is established", "12  Next confirmatory work", "Appendix  Milestone ledger and update protocol"
    ]
    add_numbered(doc, contents)

    doc.add_heading("Executive summary", level=1)
    add_table(doc, ["Area", "Current result", "Interpretation"], [
        ("Production engine", "Single-threaded C++20 UCI engine with Python oracle", "Playable in standard chess GUIs and independently testable"),
        ("Engine strength", "1996 estimate; approximate 95% interval 1930 to 2063", "A local Stockfish limited-strength benchmark, not a human-platform rating"),
        ("Adaptive safety", "MultiPV candidates, 35 cp bound, mate protection, neutral fallback", "Adaptation is constrained to objectively competitive root moves"),
        ("Personal corpus", "1,166 usable exact 10-minute games; 36,039 decisions", "Large enough for a first player-specific prediction study"),
        ("Selected model", "Histogram gradient boosting", "Best held-out probability prediction and only class where personal training beat population training"),
        ("Open claim", "No compute-matched win-rate result yet", "Predictive usefulness must still be translated into playing outcomes"),
    ], widths=[1.35, 2.45, 3.2], font_size=8.2)
    add_figure(doc, architecture, 6.9, "Figure 1  Current hybrid architecture and responsibility boundary")

    # Part one
    doc.add_heading("Part One  Main performance development", level=1)
    para(doc, "The first phase built a conventional chess engine from first principles, validated it with reference positions and regression tests, then moved performance-critical work into C++. This phase matters to the adaptive research because any claimed adaptive gain is meaningless if the neutral engine is unstable, illegal, or receives less compute than the treatment.")

    doc.add_heading("1  Foundations and correctness", level=2)
    para(doc, "The initial Python implementation established the position model: board representation, compact move encoding, FEN parsing, reversible make and unmake, castling, en passant, promotions, undo history, incremental Zobrist hashing, and internal validation. Legal chess behavior then grew through attack maps, pseudo-legal generation, legal filtering, checker and pin handling, draw rules, perft, and divide.")
    add_bullets(doc, [
        "Perft reference positions verify move counts across special rules and deeper move trees.",
        "Round-trip tests verify that every move and unmove restores the exact prior position and hash.",
        "Random-game comparisons preserve a slower correctness path as an oracle for optimized legal generation.",
        "The Python implementation remains deliberately independent after the C++ migration, allowing the production engine to be checked rather than trusted by construction."
    ])

    doc.add_heading("2  Search and evaluation", level=2)
    para(doc, "The playing engine progressed from negamax alpha-beta to iterative deepening, principal-variation search, quiescence, aspiration windows, transposition storage, search ordering, and conservative pruning. Evaluation moved from material and placement toward a tapered middlegame and endgame score with pawn structure, king shelter, rook files, bishop pair, connected and passed pawns, and piece mobility.")
    add_table(doc, ["Subsystem", "Implemented mechanisms", "Why it mattered"], [
        ("Search", "Iterative deepening, PVS, quiescence, aspiration windows", "Improved depth, move stability, and reuse of earlier iterations"),
        ("Ordering", "Hash move, MVV-LVA, SEE, killers, history, countermoves", "Raised alpha-beta cutoff efficiency"),
        ("Pruning", "LMR, null move, futility, quiescence delta and SEE pruning", "Avoided low-value work with safeguards for tactical and zugzwang cases"),
        ("Evaluation", "Tapered score, structure, shelter, rook files, mobility", "Made the neutral engine strategically less shallow"),
        ("Protocol", "UCI, clocks, increments, node limits, asynchronous stop", "Enabled repeatable GUI play and experiments"),
    ], widths=[1.05, 3.2, 2.75])

    doc.add_heading("3  Performance engineering", level=2)
    para(doc, "Optimization followed measured hot paths rather than replacing correctness checks. The move core adopted bitboard iteration, precomputed king and knight attacks, direct checker and absolute-pin legality, and occupancy-indexed sliding attacks. Search allocations were reduced through fixed-capacity move lists and principal-variation storage. Material, piece placement, bishop counts, and game phase became incremental; wider geometric terms remain calculated at leaves.")
    add_bullets(doc, [
        "Pawn structure was cached and evaluation loops became bitboard-driven.",
        "Undo-state records are reused by ply, reducing recursive allocation pressure.",
        "Static exchange evaluation improved capture ordering and prevented hopeless losing captures from expanding quiescence.",
        "Clustered transposition replacement retains several colliding positions and protects deeper current entries.",
        "Deterministic fixed-node modes made before-and-after comparisons reproducible."
    ])
    para(doc, "The engineering principle was to preserve a slower reference behavior wherever possible. Speed was accepted only with perft, tactical, search-result, and state-restoration coverage.")

    doc.add_heading("4  C++ production migration", level=2)
    para(doc, "The project became a hybrid system rather than a wholesale rewrite. C++ owns the production UCI engine, position core, move generation, evaluation, search, timing, transposition table, and root selection. Python owns the correctness oracle, experiment harnesses, PGN processing, labeling, training, uncertainty analysis, and model research.")
    add_table(doc, ["C++ production boundary", "Python research and validation boundary"], [
        ("Packed moves, bitboards, FEN and Zobrist", "Independent legal moves, perft and hash cross-checks"),
        ("Direct legal generation and lookup attacks", "Regression suites and deterministic random comparisons"),
        ("Tapered evaluation and timed PVS", "Dataset construction and reference-engine labeling"),
        ("UCI, asynchronous stop and self-play", "Model fitting, calibration, bootstrap uncertainty"),
        ("MultiPV and bounded adaptive selection", "Profile creation and future model export"),
    ], widths=[3.45, 3.55])
    para(doc, "Python is not called at every search node. The intended deployment path is to train and validate outside the engine, export a compact frozen artifact, and perform low-overhead inference in C++ only at the root candidate boundary.")

    doc.add_heading("5  Measurement and current engine baseline", level=2)
    para(doc, "The project added deterministic tactical benchmarks, fixed-node self-play, EPD best-move solving, SPRT support, and a Stockfish gauntlet. The current absolute estimate uses Stockfish 18 with UCI_LimitStrength, one thread, 64 MB hash, 30 ms per move, deterministic four-ply openings, color swaps, legal-move verification, and a 140-ply maximum.")
    add_figure(doc, elo, 6.5, "Figure 2  Measured local engine-strength progression with approximate intervals")
    add_table(doc, ["Opponent setting", "Games", "W D L", "Score", "Single-sample estimate"], [
        ("Stockfish 1875", "64", "35 14 15", "65.6%", "1987"),
        ("Stockfish 2000", "48", "20 9 19", "51.0%", "2007"),
        ("Combined fit", "112", "55 23 34", "59.4%", "1996  interval 1930 to 2063"),
    ], widths=[1.55, .7, 1.15, .9, 2.7])
    para(doc, "The earlier 78-game baseline was 1871 with an approximate interval of 1792 to 1949. The newer estimate is about 125 points higher, but it should not be interpreted as FIDE, Chess.com, Lichess, CCRL, or CEGT Elo. It is a local rating within one short-time reference protocol. All 112 latest games completed without a crash or illegal move.")
    add_status_line(doc, "Current verification", "113 Python tests passed in the project research environment on 30 September 2026. The C++ implementation also has dedicated suites for position, perft, move generation, search, pruning, timing, MultiPV, adaptive features, and adaptive selection.")

    # Part two
    doc.add_heading("Part Two  Adaptive layer and opponent modelling", level=1)
    para(doc, "The adaptive phase changed the central question from making the engine generally stronger to testing a causal claim: can prior information about a recurring opponent improve match score when the personalized and neutral engines receive identical compute and root candidates? The project first built a safe hand-authored selector, then created a data-driven personal error-prediction study.")

    doc.add_heading("6  Research question and safety contract", level=2)
    para(doc, "The treatment is not allowed to search more moves or spend more time than the control. Both arms use MultiPV 4 and the same all-root search. Neutral play selects rank one; adaptive play may rerank only completed candidates. Single-PV neutral play is a practical baseline but not the causal control because it performs less root work.")
    add_table(doc, ["Invariant", "Operational rule"], [
        ("Legality", "Only moves returned by legal C++ root search can be selected"),
        ("Objective safety", "Alternatives must be no more than 35 centipawns below the neutral best move"),
        ("Mate safety", "Winning mates are preserved and lines allowing forced mate are ineligible"),
        ("Completeness", "Only a fully completed all-root MultiPV iteration can drive adaptation"),
        ("Fallback", "Invalid profiles, interrupted search, terminal positions, or unsupported data fail closed to neutral"),
        ("Reproducibility", "Fixed position, profile, engine version, limits, and seed must reproduce the decision"),
        ("Fair experiment", "Binary, nodes, openings, colors, hardware, and prior-information boundary are matched"),
    ], widths=[1.35, 5.65])

    doc.add_heading("7  Hand-authored adaptive layer", level=2)
    doc.add_heading("7.1 MultiPV root candidates", level=3)
    para(doc, "The C++ engine was extended to return multiple unique, score-sorted root candidates together with completion metadata. Terminal roots are represented as complete zero-candidate results. Interrupted searches can retain diagnostics but cannot become an adaptive decision set.")
    doc.add_heading("7.2 Position characteristics", level=3)
    para(doc, "Every completed root candidate is temporarily made on the board and described from the root mover's perspective. The extractor restores the original position. These values measure the position and move; they do not themselves contain preferences.")
    add_table(doc, ["Feature group", "Examples"], [
        ("Forcing play", "capture value, promotion gain, check, legal reply count"),
        ("Material", "balance, absolute imbalance, remaining non-pawn material, pawn count"),
        ("Position shape", "open files, mobility, king safety, pawn structure, center control"),
        ("Commitment", "pawn tension, irreversible move, castling"),
    ], widths=[1.8, 5.2])
    doc.add_heading("7.3 Versioned profiles and synthetic archetypes", level=3)
    para(doc, "Opponent profiles use a strict JSON schema with a signed weight and confidence for every feature. Unknown fields, duplicates, unsupported versions, missing features, non-integer values, non-finite values, and out-of-range values are rejected. The profile cannot change engine-owned safety limits.")
    add_table(doc, ["Synthetic profile", "Intended pressure"], [
        ("Tactical pressure", "Checks, captures, restricted replies, open and imbalanced play"),
        ("Simplification pressure", "Exchanges, reduced material and tension, safer conversion"),
        ("Complexity pressure", "More pieces, choices, tension, imbalance, and commitment"),
        ("Positional restriction", "Mobility, king safety, structure, center control, restricted replies"),
    ], widths=[2.1, 4.9])
    doc.add_heading("7.4 Bounded root selector", level=3)
    para(doc, "Candidate zero remains the neutral move. Eligible alternatives are compared with it using normalized feature differences, profile weights, and profile confidence. Only a strictly positive adaptive utility can change the move; ties preserve neutral order. Adaptive Mode is opt-in and the profile is copied into the active search task, preventing mid-search mutation.")
    para(doc, "This implementation proves that the engine can express controlled behavioral differences. It does not prove that any synthetic style exploits a real person, and it does not yet emit the full candidate-by-candidate decision explanation originally planned.")

    doc.add_heading("8  Personal-data research pipeline", level=2)
    add_figure(doc, chain, 7.0, "Figure 3  Evidence pipeline from public games to the future playing experiment")
    doc.add_heading("8.1 Corpus", level=3)
    para(doc, "The read-only Chess.com archive pipeline downloaded 1,676 public games for YashDutt7. Of these, 1,182 used the exact 600-second time control. The parser produced 36,039 target-player decisions from 1,166 usable games. Raw data and fitted artifacts are ignored by Git; code, schemas, hashes, aggregate results, and protocols are versioned.")
    add_table(doc, ["Corpus quantity", "Value"], [
        ("All public games downloaded", "1,676"), ("Exact 10-minute games", "1,182"),
        ("Usable parsed games", "1,166"), ("Personal decisions", "36,039"),
        ("Eligible non-mate labels", "32,988"), ("Large errors at 100 cp", "5,726"),
        ("Mate-score exclusions", "3,051"), ("Archive PGN SHA-256", "d624add4...e4e2a6f"),
    ], widths=[3.2, 3.8])
    doc.add_heading("8.2 Chronological split", level=3)
    para(doc, "Whole games were sorted by date and divided 60 percent training, 20 percent validation, and 20 percent final test. Positions from the same game never cross partitions. This simulates actual deployment: only past information can predict later behavior.")
    add_table(doc, ["Split", "Dates", "Games", "Decisions", "Mean rating", "Range"], [
        ("Train", "2022-12-16 to 2024-06-12", "695", "20,320", "772.5", "370 to 950"),
        ("Validation", "2024-06-13 to 2025-01-21", "235", "7,735", "985.1", "897 to 1,095"),
        ("Test", "2025-01-21 to 2026-06-27", "236", "7,984", "1,186.1", "1,056 to 1,296"),
    ], widths=[.8, 2.0, .65, .85, .9, 1.35], font_size=7.8)
    doc.add_heading("8.3 Labels and stability", level=3)
    para(doc, "For each decision, Stockfish first evaluated the original position and then the child position after the played move. Scores were converted to the player's pre-move perspective. A large error is a loss of at least 100 centipawns; mate-score cases are excluded because mate distance is not a stable centipawn quantity.")
    para(doc, "The node budget was locked before fitting through a deterministic 500-position stability study. Agreement improved monotonically with more nodes. The chosen 100,000-node budget still contains measurement noise: 12 of 453 jointly eligible labels changed between 50,000 and 100,000 nodes.")
    add_table(doc, ["Budgets", "Agreement", "Kappa", "Flips", "CPL correlation"], [
        ("5k versus 20k", "94.44%", "0.7980", "26", "0.9549"),
        ("20k versus 50k", "96.73%", "0.8859", "15", "0.9745"),
        ("50k versus 100k", "97.35%", "0.9071", "12", "0.9827"),
    ], widths=[1.8, 1.2, 1.0, .8, 2.2])
    doc.add_heading("8.4 Leakage-resistant model features", level=3)
    para(doc, "The model sees only information available before the player's decision: check state, legal moves and captures, material, remaining material, pawns, open files, mobility, king safety, pawn structure, center control, tension, move number, color, player rating, and rating difference. It does not receive Stockfish evaluation, the played move's outcome, centipawn loss, game result, or future information.")

    doc.add_heading("9  Three-model comparison", level=2)
    para(doc, "Three model families were compared because they represent increasing flexibility and different assumptions about personalization.")
    add_table(doc, ["Model", "What it assumes", "Role in the study"], [
        ("Regularized logistic regression", "One mostly additive relationship between features and error probability", "Transparent, calibrated linear baseline"),
        ("Hierarchical Bayesian logistic", "Population coefficients plus a cautious player-specific deviation", "Partial pooling and uncertainty, implemented with a Laplace approximation"),
        ("Histogram gradient boosting", "Conditional thresholds and interactions among position features", "Nonlinear challenger able to represent different regimes"),
    ], widths=[2.0, 2.8, 2.2])
    add_figure(doc, models, 6.7, "Figure 4  Personalized model test log loss on untouched chronological games")
    add_table(doc, ["Personalized model", "Log loss", "Brier", "ROC AUC", "Avg precision", "ECE 10"], [
        ("Logistic", "0.40499", "0.12157", "0.62550", "0.21108", "0.03020"),
        ("Hierarchical logistic", "0.40471", "0.12152", "0.62584", "0.21121", "0.02968"),
        ("Gradient boosting", "0.39348", "0.11958", "0.64934", "0.22044", "0.01608"),
    ], widths=[2.0, 1.0, 1.0, 1.0, 1.1, .9], font_size=8.0)
    para(doc, "Lower log loss and Brier score mean better probability predictions. ROC AUC measures ranking ability: the boosting value of 0.649 means that, for a randomly selected error and non-error example, the model ranks the error as more likely about 65 percent of the time. Average precision focuses on the rarer positive class, while ECE summarizes calibration error across probability bins.")
    add_table(doc, ["Model class", "Personal", "Population", "Personal minus population", "95% game-bootstrap interval", "Finding"], [
        ("Logistic", "0.40499", "0.40299", "+0.00200", "+0.00025 to +0.00368", "Population better"),
        ("Hierarchical", "0.40471", "0.40291", "+0.00180", "+0.00038 to +0.00313", "Population better"),
        ("Gradient boosting", "0.39348", "0.40004", "-0.00656", "-0.01081 to -0.00198", "Personal better"),
    ], widths=[1.25, .8, .85, 1.15, 1.85, 1.1], font_size=7.4)
    para(doc, "The central result is nuanced. Personal history did not automatically help. Both linear personal models were slightly worse than population versions on the same future Yash positions. Only nonlinear gradient boosting converted the personal history into a statistically supported improvement. A 2,000-replicate paired bootstrap over whole games estimated a 99.9 percent probability that personalized boosting was better than population boosting on this test.")
    para(doc, "Post-hoc ablations removed rating and time proxies. The personal boosting advantage persisted without player rating and rating difference, and also without rating plus fullmove number. These checks are consistent with nonlinear position-dependent personal tendencies, but they are exploratory because they were run after inspecting the primary result.")

    doc.add_heading("10  Player growth and concept drift", level=2)
    add_figure(doc, drift, 6.6, "Figure 5  Rating drift across chronological partitions; vertical lines show observed ranges")
    para(doc, "The training and test periods do not describe a stationary player. Mean rating increased by about 414 points, and the maximum training rating of 950 is below the minimum test rating of 1,056. A lifetime profile therefore combines early, intermediate, and later versions of the player.")
    para(doc, "Chronological testing was still the correct primary choice because a random split would leak later-strength behavior into training and answer an easier question. The cost is that the experiment tests forward personalization under improvement rather than interpolation within a stable player state.")
    para(doc, "The nonlinear model did not eliminate growth. It handled the mixture better by learning conditional regimes: a feature can matter differently in a tense middlegame than in a quiet ending, or at different ratings and material structures. Tree models still cannot smoothly extrapolate beyond observed rating thresholds, and the selected model probably retains obsolete information.")
    para(doc, "The population data are also heterogeneous, but the population model estimates general relationships between position types and human error rather than one fictional combined person. Variation can average into a stable baseline. Personalization must learn tendencies that are specific to the player and remain present in future play, which is a stricter requirement.")

    doc.add_heading("11  What is established", level=2)
    add_table(doc, ["Established by current evidence", "Not yet established"], [
        ("The production engine is legal, testable, GUI-compatible, and locally measured", "The 1996 estimate transfers to human or public rating pools"),
        ("Position features predict future large errors better than a constant rate", "The model explains causal psychological weaknesses"),
        ("Boosting outperforms both linear alternatives on future games", "The result transfers to other players or time controls"),
        ("Personal boosting beats same-class population boosting on Yash test games", "Personalized root selection wins more games at equal compute"),
        ("Linear personalization can be counterproductive under strong player drift", "A lifetime profile is the best way to deploy personalization"),
    ], widths=[3.5, 3.5])

    doc.add_heading("12  Next confirmatory work", level=2)
    para(doc, "The next phase must translate probability prediction into a frozen playing policy without tuning on the final match set.")
    add_numbered(doc, [
        "Freeze and export the selected boosting model with an exact feature order, missing-value policy, thresholds, leaf values, and a Python-to-C++ parity test corpus.",
        "Implement C++ inference after each safe MultiPV candidate. Predict the opponent's next-move large-error probability from the resulting position.",
        "Define and preregister a bounded probability bonus. The 35 cp objective eligibility bound and all mate protections remain authoritative.",
        "Run identical-node, identical-opening, color-swapped comparisons among neutral, population, personalized, wrong-player, and random-safe selectors.",
        "Use paired game-level uncertainty or SPRT. Do not tune the bonus on the confirmatory match set.",
        "After the profile and policy are frozen, run prospective blinded human games. Retrospective prediction cannot substitute for this causal result.",
        "In parallel, compare lifetime, recent-window, time-decayed, and online-updated profiles chronologically to model the player as a changing state rather than a permanent vector."
    ])
    para(doc, "The strongest research direction is opponent modelling under concept drift: whether an engine can track an evolving human well enough to improve compute-matched playing outcomes. This framing preserves negative findings and makes player growth part of the scientific problem instead of treating it as inconvenient noise.")

    doc.add_heading("Appendix  Milestone ledger", level=1)
    add_table(doc, ["Date", "Milestone", "Evidence retained"], [
        ("10 Jul 2026", "Core board representation and move system", "Stage 1 implementation and regression tests"),
        ("28 Jul 2026", "Legal generation, evaluation, alpha-beta, UCI, TT, ordering", "Perft, search and protocol tests"),
        ("4 Aug 2026", "Python performance work and C++ production migration", "Benchmarks, cross-checks, self-play and SPRT tools"),
        ("18 Aug 2026", "C++ direct legality, sliding lookups, preallocation, incremental evaluation", "Architecture sync and 1996 Elo baseline"),
        ("18 Aug to 1 Sep 2026", "Adaptive contract, MultiPV, features, schema, synthetic profiles, selector", "Versioned docs, profiles and C++ tests"),
        ("30 Sep 2026", "Personal PGN pipeline, reference labels, features and three-model comparison", "Dataset hashes, stability report, drift report and held-out results"),
        ("30 Sep 2026", "Living project report established", "This Word report and reproducible builder"),
    ], widths=[1.25, 3.3, 2.45], font_size=8.0)

    doc.add_heading("Report maintenance rule", level=2)
    para(doc, "This document is a living project artifact. Every future GitHub push that changes architecture, experimental protocol, data, model behavior, measurements, conclusions, or the roadmap should update the corresponding section of this report in the same work cycle. Small refactors that do not change behavior may be recorded in the milestone ledger rather than expanded into a new section.")
    add_bullets(doc, [
        "Record the commit or milestone, date, purpose, implementation change, validation performed, result, and remaining uncertainty.",
        "Separate measured facts from interpretations and proposed next steps.",
        "Preserve negative results and failed hypotheses; do not rewrite history around the winning model.",
        "Update charts and sample counts when a newer experiment supersedes them, while retaining the older value in the milestone narrative.",
        "Regenerate this Word file with tools/build_project_report.py, render it, inspect every page, then commit the report with the related project change."
    ])

    doc.add_heading("Source record", level=2)
    para(doc, "This version was assembled from the Git history through commit f3cd286, the repository roadmap and architecture notes, the Elo benchmark, the adaptive experiment contract, the profile and feature specifications, the reference-label stability study, the player-drift report, and the YashDutt7 model-results report. Personal raw games and fitted model artifacts remain outside version control by design.")

    core = doc.core_properties
    core.title = "Adaptive Chess Engine Project Development and Research Report"
    core.subject = "Living technical report covering engine performance and adaptive opponent modelling"
    core.author = "Yash Dutt"
    core.keywords = "chess engine, adaptive engine, opponent modelling, gradient boosting, C++, research"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build_report()
    print(path)
